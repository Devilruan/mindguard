# -*- coding: utf-8 -*-
"""生成《作品简介.docx》- 赛题5版，≤300字"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import os

OUT = r'H:\2026AIC·算法创新赛\提交材料_赛题5-AI学科交叉\5-作品简介\作品简介_神经符号引擎驱动的校园心理风险智能筛查与干预系统.docx'

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(3.0)

def add_p(text, size=12, bold=False, cn='宋体', align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = 'Times New Roman'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    return p

add_p('作品简介', 16, True, '黑体')
add_p('')
add_p('（不超过300字）', 10, False, '楷体')
add_p('')

intro = (
    '【问题】校园心理筛查面临单一量表漏检高（70%隐性风险在PHQ-9上<10分）、'
    'AI黑箱不可解释、建议同质化三大痛点。\n'
    '【方案】立足精神医学×人工智能交叉，提出神经符号双引擎筛查系统：'
    '融合PHQ-9/GAD-7临床量表与开放文本，经规则知识库（6条DSM-5规则+RIPPER归纳）'
    '与4级危机分级拦截，量表正常时仍能检出隐性危机信号，RAG知识库自动生成个性化干预建议。\n'
    '【效果】118例对比实验：额外检出12例纯量表漏检的隐性危机，'
    '检出率净提升16.3%、假阳性减72.7%，每条结论附完整规则激活链。\n'
    '【应用】纯Python实现，PyInstaller一键打包，可部署至任何高校心理中心。'
)
add_p(intro, 12)
add_p('')
# 统计字数
import re
chars = len(re.sub(r'\s', '', intro))
add_p(f'（全文{chars}字）', 10, False, '楷体')

doc.save(OUT)
print(f'✅ 作品简介已生成: {OUT} ({chars}字)')