# -*- coding: utf-8 -*-
"""
backend.api.export_api
报告导出：生成 Word (.docx) 筛查报告（含风险分、推理链、规则详情、心理热线）。
"""
from __future__ import annotations
import io
import json
import logging
from typing import Optional

from fastapi import APIRouter, HTTPException, Depends, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from backend.models.database import ScreeningSession, get_db
from config.settings import APP_NAME

log = logging.getLogger(__name__)
router = APIRouter()


# Admin 路由保护依赖（引用共享模块，避免 main ↔ api 循环导入）
from config.admin_auth import require_admin_auth as _require_admin_auth


def _build_docx(session: ScreeningSession, session_dict: dict) -> bytes:
    """
    用 python-docx 构建一份完整筛查报告，返回 .docx 的二进制内容。
    """
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()

    # 默认中文字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is not None:
        rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # ── 标题 ──
    title = doc.add_heading(f'{APP_NAME}', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('心理健康风险筛查报告')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor(70, 70, 70)

    doc.add_paragraph()

    # ── 基本信息 ──
    doc.add_heading('基本信息', level=1)
    created = session.created_at.strftime('%Y-%m-%d %H:%M:%S') if session.created_at else '-'
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Light Grid Accent 1'
    cells = table.rows[0].cells
    cells[0].text = '筛查编号'
    cells[1].text = f'#{session.id}'
    cells = table.rows[1].cells
    cells[0].text = '时间'
    cells[1].text = created
    cells = table.rows[2].cells
    cells[0].text = '作答时长'
    cells[1].text = f'{session.duration_sec or 0:.0f} 秒'
    cells = table.rows[3].cells
    cells[0].text = '匿名标识'
    u = session.user
    cells[1].text = u.anon_id if u else '未知'

    doc.add_paragraph()

    # ── 风险总分 ──
    doc.add_heading('风险评估结果', level=1)

    level = session.overall_risk_level or 'low'
    level_cn = {'low': '低风险', 'medium': '中风险', 'high': '高风险', 'urgent': '极高风险'}
    level_display = level_cn.get(level, level)

    p = doc.add_paragraph()
    run = p.add_run(f'综合风险等级：{level_display}')
    run.bold = True
    run.font.size = Pt(14)
    # 不同等级用不同颜色
    level_colors = {
        'low': RGBColor(46, 160, 67),
        'medium': RGBColor(217, 172, 12),
        'high': RGBColor(215, 58, 73),
        'urgent': RGBColor(163, 14, 21),
    }
    run.font.color.rgb = level_colors.get(level, RGBColor(70, 70, 70))

    doc.add_paragraph(f'综合风险得分：{session.overall_risk_score:.1f} / 100')
    doc.add_paragraph(f'抑郁维度：{session.depression_score:.1f} / 100  （PHQ-9={session.phq9_score}）')
    doc.add_paragraph(f'焦虑维度：{session.anxiety_score:.1f} / 100  （GAD-7={session.gad7_score}）')
    doc.add_paragraph(f'自杀风险：{session.suicide_risk_score:.1f} / 100')

    doc.add_paragraph()

    # ── 激活规则 ──
    doc.add_heading('推理依据', level=1)
    activated = session.activated_rules
    if isinstance(activated, str):
        try:
            activated = json.loads(activated)
        except Exception:
            activated = []

    if activated:
        for r in activated:
            rid = r.get('rule_id', '?')
            explain = r.get('explain', '')
            bullet = doc.add_paragraph(style='List Bullet')
            run_id = bullet.add_run(f'[{rid}] ')
            run_id.bold = True
            bullet.add_run(explain)
    else:
        doc.add_paragraph('本次筛查未触发任何高优先级规则。')

    doc.add_paragraph()

    # ── 推理链路 ──
    doc.add_heading('完整推理链路', level=1)
    chain = session.explanation_chain
    if isinstance(chain, str):
        try:
            chain = json.loads(chain)
        except Exception:
            chain = []
    if chain:
        for step in chain:
            p = doc.add_paragraph(step)
            p.paragraph_format.left_indent = Cm(0.5)
    else:
        doc.add_paragraph('（无）')

    doc.add_paragraph()

    # ── 心理资源 ──
    doc.add_heading('心理自助资源', level=1)
    doc.add_paragraph('全国24小时心理援助热线：12356')
    doc.add_paragraph('北京市心理危机研究与干预中心：010-82951332（24小时）')
    doc.add_paragraph('青少年心理咨询热线：12355（工作日 9:00-21:00）')
    doc.add_paragraph()
    doc.add_paragraph('推荐自我调节技巧：')
    tips = [
        '4-7-8 呼吸法：吸气 4 秒，屏息 7 秒，缓慢呼气 8 秒，重复 4 轮',
        '着陆技术：说出你看到的 5 样东西 → 3 种声音 → 1 种触摸感受',
        '渐进式肌肉放松：从脚趾开始，依次收紧-放松全身肌肉群',
    ]
    for t in tips:
        doc.add_paragraph(t, style='List Bullet')

    doc.add_paragraph()
    doc.add_paragraph('本报告由神经符号心理健康风险智能筛查系统生成，仅供参考，不构成医学诊断。')

    # 输出到字节流
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get('/session/{session_id}/report.docx')
def export_report(session_id: int, _: bool = Depends(_require_admin_auth), db: Session = Depends(get_db)):
    """
    生成指定筛查会话的 Word 报告并下载。
    """
    session = db.query(ScreeningSession).filter(ScreeningSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail='会话不存在')

    # 避免循环 SQLAlchemy 警告
    session_dict = {
        'activated_rules': session.activated_rules,
        'explanation_chain': session.explanation_chain,
    }

    try:
        data = _build_docx(session, session_dict)
    except ImportError:
        raise HTTPException(status_code=500, detail='服务器缺少 python-docx 依赖')
    except Exception as e:
        log.exception('生成 Word 报告失败')
        raise HTTPException(status_code=500, detail=f'报告生成失败: {e}')

    level = session.overall_risk_level or 'unknown'
    created = session.created_at.strftime('%Y%m%d_%H%M%S') if session.created_at else 'unknown'
    filename = f'mmhs_report_{session_id}_{level}_{created}.docx'

    log.info('导出 Word 报告 session=%d size=%d bytes', session_id, len(data))

    return StreamingResponse(
        io.BytesIO(data),
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )

